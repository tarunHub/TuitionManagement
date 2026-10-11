#!/usr/bin/env python3
"""Replace the portal's data with the contents of SmartStart_Student_Management.xlsx.

Prints one SQL transaction (wipe + inserts) on stdout and a report on stderr; nothing touches the database
until you pipe it into psql:

    python3 tools/import_excel.py src/main/resources/SmartStart_Student_Management.xlsx > import.sql
    docker compose exec -T postgres sh -c 'psql -v ON_ERROR_STOP=1 -U "$POSTGRES_USER" -d "$POSTGRES_DB"' < import.sql

Kept: admin logins and notification templates. Everything else (students, groups, grade rates, attendance,
fees, payments, feedback, notification history, terms, schedule slots) is replaced.
Requires openpyxl (pip install openpyxl).

How the workbook is read:
- Students sheet is the source of truth. Rows without a child's name (and "(Example)" rows) are skipped.
  "Inquiry" becomes INACTIVE. The "Alternate Contact" column holds the batch (Batch1..4) in practice, so a
  BatchN value there becomes the class group; a second number in "Primary Contact" becomes the alternate contact.
- Monthly sheets number students by row, not by their real Student ID, so they are matched by first name.
- Attendance: P/A only ("S" = scheduled future class, other marks are reported and skipped).
- Fees: amount due = classes held (from attendance) x rate, like the app; the sheet's classes-scheduled value is
  kept as a manual override. Months with no attendance and no payments (Nov/Dec placeholders) are skipped.
- One payment per non-zero "Amount Paid". Missing dates become the month's last day (not after today), payment
  dates more than a year off the fee month are assumed to be year typos, and missing methods become OTHER.
"""
import calendar
import datetime as dt
import re
import sys
from decimal import Decimal

import openpyxl

MONTHS = {m: i for i, m in enumerate(calendar.month_abbr) if m}
PHONE = re.compile(r"^\+?[0-9][0-9 ().-]{5,18}[0-9]$")
BATCH = re.compile(r"^batch\s*\d+$", re.I)
METHODS = {"cash": "CASH", "e-transfer": "E_TRANSFER", "etransfer": "E_TRANSFER", "interac": "E_TRANSFER",
           "cheque": "CHEQUE", "check": "CHEQUE", "card": "CARD"}
TODAY = dt.date.today()

notes = []


def report(msg):
    notes.append(msg)


def text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def q(v):
    """SQL literal."""
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, Decimal)):
        return str(v)
    if isinstance(v, dt.date):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return dt.date.fromisoformat(str(v).strip()[:10])


def money(v):
    return Decimal(str(v or 0)).quantize(Decimal("0.01"))


def key(name):
    """Match key for names across sheets: first word, lower case ("Gurshabd Singh Rawala" -> "gurshabd")."""
    return str(name).strip().split()[0].lower()


def sheet_month(title):
    m = re.match(r"^(Attendance|Fees|Report)\s+([A-Za-z]{3})\s+(\d{4})$", title.strip())
    return (m.group(1), dt.date(int(m.group(3)), MONTHS[m.group(2).title()], 1)) if m else (None, None)


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def header_index(table, first_cell):
    for i, r in enumerate(table):
        if r and text(r[0]) == first_cell:
            return i
    raise ValueError(f"no header row starting with {first_cell!r}")


def col(header, *needles):
    for i, h in enumerate(header):
        h = (text(h) or "").lower()
        if all(n in h for n in needles):
            return i
    return None


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sql = ["BEGIN;",
           "TRUNCATE attendance, fee_payment, monthly_fee, monthly_feedback, notification, schedule_slot, term,",
           "         student_subject, student, class_group, grade_rate RESTART IDENTITY CASCADE;"]

    # ---- grade rates
    fs = rows(wb["Fee Structure"])
    grades = {}
    for r in fs[header_index(fs, "Grade") + 1:]:
        g, rate = text(r[0]), r[1]
        if g and isinstance(rate, (int, float)) and rate > 0:
            grades[g] = money(rate)
            sql.append(f"INSERT INTO grade_rate (grade, per_class_rate, notes) VALUES ({q(g)}, {q(money(rate))}, {q(text(r[2]))});")

    # ---- students
    st = rows(wb["Students"])
    h = st[header_index(st, "Student ID")]
    c = {name: col(h, *needles) for name, needles in {
        "code": ("student id",), "name": ("full name",), "dob": ("date of birth",), "grade": ("grade",),
        "school": ("school",), "subjects": ("subjects",), "address": ("address",), "guardian": ("guardian",),
        "primary": ("primary contact",), "alternate": ("alternate",), "email": ("email",),
        "ename": ("emergency contact name",), "ephone": ("emergency contact phone",), "enrolled": ("enrollment",),
        "rate": ("rate",), "status": ("status",), "notes": ("notes",)}.items()}
    students, groups, max_code = {}, [], 0
    for r in st[header_index(st, "Student ID") + 1:]:
        code, name = text(r[c["code"]]), text(r[c["name"]])
        if not code or not name or name.startswith("(Example)"):
            if code and not (name or "").startswith("(Example)") and any(text(v) for v in r[1:]):
                report(f"Skipped {code}: no child name (guardian {text(r[c['guardian']]) or '-'})")
            continue
        grade = text(r[c["grade"]])
        if grade not in grades:
            raise SystemExit(f"{code} {name}: grade {grade!r} is not in Fee Structure")
        phones = [p.strip() for p in str(r[c["primary"]] or "").splitlines() if p.strip()]
        alt = text(r[c["alternate"]])
        group = None
        if alt and BATCH.match(alt):
            group, alt = alt.replace(" ", "").title(), None
        if len(phones) > 1:
            alt = alt or phones[1]
            if len(phones) > 2:
                report(f"{code} {name}: extra phone numbers dropped: {', '.join(phones[2:])}")
        primary = phones[0] if phones else None
        for label, p in (("primary", primary), ("alternate", alt)):
            if p and not PHONE.match(p):
                report(f"{code} {name}: {label} contact {p!r} doesn't match the app's phone format; fix it when editing")
        raw_status = (text(r[c["status"]]) or "Active").lower()
        status = "ACTIVE" if raw_status == "active" else "INACTIVE"
        note = text(r[c["notes"]])
        if raw_status not in ("active", "inactive"):
            note = f"Status in spreadsheet: {text(r[c['status']])}" + (f". {note}" if note else "")
            report(f"{code} {name}: status {text(r[c['status']])!r} imported as INACTIVE")
        dob = to_date(r[c["dob"]])
        if dob and grade.startswith("Grade") and (dt.date(2026, 9, 1) - dob).days < 365 * 5:
            report(f"{code} {name}: date of birth {dob} looks too recent for {grade}; please check")
        rate = money(r[c["rate"]]) if r[c["rate"]] not in (None, "") else grades[grade]
        if group and group not in groups:
            groups.append(group)
        k = key(name)
        if k in students:
            raise SystemExit(f"Two students share the first name {k!r}; name matching would be ambiguous")
        students[k] = code
        max_code = max(max_code, int(re.sub(r"\D", "", code)))
        email = text(r[c["email"]])
        sql.append(
            "INSERT INTO student (student_code, full_name, date_of_birth, grade, school, address, guardian_name,"
            " primary_contact, alternate_contact, email, emergency_contact_name, emergency_contact_phone,"
            " enrollment_date, per_class_rate, status, class_group_id, notes, sms_opt_in) VALUES ("
            + ", ".join([q(code), q(name), q(dob), q(grade), q(text(r[c["school"]])), q(text(r[c["address"]])),
                         q(text(r[c["guardian"]]) or "Unknown"), q(primary), q(alt), q(email.lower() if email else None),
                         q(text(r[c["ename"]])), q(text(r[c["ephone"]])), q(to_date(r[c["enrolled"]])), q(rate),
                         q(status), f"(SELECT id FROM class_group WHERE name = {q(group)})" if group else "NULL",
                         q(note), "FALSE"]) + ");")
        for subject in [s.strip() for s in str(r[c["subjects"]] or "").split(",") if s.strip()]:
            sql.append(f"INSERT INTO student_subject (student_id, subject) SELECT id, {q(subject)} FROM student WHERE student_code = {q(code)};")
    # groups must exist before students reference them
    first_student = next(i for i, s in enumerate(sql) if s.startswith("INSERT INTO student "))
    sql[first_student:first_student] = [f"INSERT INTO class_group (name) VALUES ({q(g)});" for g in sorted(groups)]
    sql.append(f"SELECT setval('student_code_seq', {max_code});")

    def student_ref(name, where):
        k = key(name)
        if k not in students:
            report(f"{where}: '{name}' is not in the Students sheet; row skipped")
            return None
        return f"(SELECT id FROM student WHERE student_code = {q(students[k])})"

    # ---- attendance (also gives classes held per student-month for fees)
    held = {}
    for ws in wb.worksheets:
        kind, month = sheet_month(ws.title)
        if kind != "Attendance":
            continue
        t = rows(ws)
        hi = header_index(t, "Student ID")
        days = {i: int(v) for i, v in enumerate(t[hi]) if isinstance(v, (int, float)) and 1 <= v <= 31}
        for r in t[hi + 1:]:
            if not text(r[0]) or not text(r[1]):
                continue
            ref = student_ref(r[1], ws.title)
            if not ref:
                continue
            for i, day in days.items():
                mark = (text(r[i]) or "").upper()
                if mark in ("P", "A"):
                    sql.append(f"INSERT INTO attendance (student_id, class_date, status) VALUES ({ref}, "
                               f"{q(month.replace(day=day))}, {q('PRESENT' if mark == 'P' else 'ABSENT')});")
                    held[(key(r[1]), month)] = held.get((key(r[1]), month), 0) + 1
                elif mark and mark != "S":
                    report(f"{ws.title}: {r[1]} on day {day} is marked {mark!r} (not P/A); skipped")

    # ---- fees and payments
    for ws in wb.worksheets:
        kind, month = sheet_month(ws.title)
        if kind != "Fees":
            continue
        t = rows(ws)
        hi = header_index(t, "Student ID")
        h = t[hi]
        fc = {"rate": col(h, "rate"), "held": col(h, "held"), "sched": col(h, "schedul"), "paid": col(h, "paid"),
              "date": col(h, "payment date"), "method": col(h, "method"), "notes": col(h, "notes")}
        body = [r for r in t[hi + 1:] if text(r[0]) and text(r[1])]
        if not any(held.get((key(r[1]), month)) or money(r[fc["paid"]]) > 0 for r in body):
            report(f"{ws.title}: no attendance or payments yet; skipped (use Generate fees in the app)")
            continue
        last_day = month.replace(day=calendar.monthrange(month.year, month.month)[1])
        for r in body:
            ref = student_ref(r[1], ws.title)
            if not ref:
                continue
            rate = money(r[fc["rate"]])
            n = held.get((key(r[1]), month), 0)
            sheet_held = int(r[fc["held"]] or 0) if fc["held"] is not None else n
            if sheet_held != n:
                report(f"{ws.title}: {r[1]} sheet says {sheet_held} classes held, attendance has {n}; using {n}")
            sched = int(r[fc["sched"]]) if fc["sched"] is not None and isinstance(r[fc["sched"]], (int, float)) else None
            note = text(r[fc["notes"]])
            note = None if note and BATCH.match(note) else note
            sql.append(
                "INSERT INTO monthly_fee (student_id, fee_month, rate, classes_held, classes_scheduled, scheduled_override,"
                f" amount_due, notes) VALUES ({ref}, {q(month)}, {q(rate)}, {n}, {q(sched)}, {q(sched is not None)},"
                f" {q(rate * n)}, {q(note)});")
            paid = money(r[fc["paid"]])
            if paid <= 0:
                continue
            ref_note = ["Excel import"]
            pdate = to_date(r[fc["date"]]) if fc["date"] is not None else None
            if pdate is None:
                pdate = min(last_day, TODAY)
                ref_note.append("date not recorded")
            elif abs((pdate - month).days) > 366:
                fixed = pdate.replace(year=month.year)
                report(f"{ws.title}: {r[1]} payment date {pdate} looks like a typo; using {fixed}")
                pdate = fixed
            raw_method = (text(r[fc["method"]]) if fc["method"] is not None else None) or ""
            method = METHODS.get(raw_method.lower().replace(" ", ""), "OTHER")
            if method == "OTHER":
                ref_note.append("method not recorded" if not raw_method else f"method: {raw_method}")
            sql.append(
                "INSERT INTO fee_payment (monthly_fee_id, amount, payment_date, method, reference) VALUES ("
                f"(SELECT id FROM monthly_fee WHERE student_id = {ref} AND fee_month = {q(month)}),"
                f" {q(paid)}, {q(pdate)}, {q(method)}, {q('; '.join(ref_note))});")

    # ---- monthly feedback
    for ws in wb.worksheets:
        kind, month = sheet_month(ws.title)
        if kind != "Report":
            continue
        t = rows(ws)
        for r in t[header_index(t, "Student ID") + 1:]:
            if not text(r[0]) or not text(r[1]) or not any(text(v) for v in r[2:5]):
                continue
            ref = student_ref(r[1], ws.title)
            if ref:
                sql.append("INSERT INTO monthly_feedback (student_id, feedback_month, topics_taught, improvement_areas, achievements)"
                           f" VALUES ({ref}, {q(month)}, {q(text(r[2]) or '')}, {q(text(r[3]))}, {q(text(r[4]))});")

    # ---- term (schedule rows in the sheet are examples only)
    sc = rows(wb["Schedule"])
    cfg = {text(r[0]): r[1] for r in sc if r and text(r[0])}
    start, end = to_date(cfg.get("Term Start")), to_date(cfg.get("Term End"))
    if start and end:
        season = "Fall" if start.month >= 8 else "Winter" if start.month <= 3 else "Spring"
        sql.append(f"INSERT INTO term (name, start_date, end_date, calendar_name, active) VALUES "
                   f"({q(f'{season} {start.year}')}, {q(start)}, {q(end)}, {q(text(cfg.get('Calendar Name')))}, TRUE);")
    hi = header_index(sc, "Group / Class")
    examples = [r for r in sc[hi + 1:] if text(r[0]) and text(r[0]).startswith("(Example)")]
    real = [r for r in sc[hi + 1:] if text(r[0]) and not text(r[0]).startswith("(Example)") and text(r[1])]
    if examples:
        report(f"Schedule: {len(examples)} example class slot(s) skipped; add real slots on the Schedule page")
    if real:
        report(f"Schedule: {len(real)} non-example slot row(s) found but not imported; add them on the Schedule page")

    log = rows(wb["Payments Log"])
    if any(r for r in log[1:] if text(r[0]) and "(Example)" not in str(r)):
        report("Payments Log has non-example rows; they were NOT imported (payments come from the Fees sheets)")

    sql.append("COMMIT;")
    print("\n".join(sql))
    print(f"Students: {len(students)}  Groups: {', '.join(sorted(groups))}  Grade rates: {len(grades)}", file=sys.stderr)
    for n in notes:
        print(" - " + n, file=sys.stderr)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
